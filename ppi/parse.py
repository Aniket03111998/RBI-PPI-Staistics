"""XLSX -> tidy long dataframe. Column positions drift across months (14 vs 16
metric columns seen in Phase 0), so headers are read from row text, never fixed indices."""
from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

from ppi.config import ALIASES_PATH

MISSING_MARKERS = {"", "-", "NA", "N.A.", "N/A", "na"}


def _clean_text(v) -> str:
    if v is None:
        return ""
    s = "".join(ch for ch in str(v) if unicodedata.category(ch) not in ("Cc", "Co", "Cf"))
    return re.sub(r"\s+", " ", s).strip()


def coerce_number(v) -> float:
    if v is None:
        return float("nan")
    if isinstance(v, (int, float)):
        return float(v)
    s = _clean_text(v).replace(",", "")
    if s in MISSING_MARKERS:
        return float("nan")
    try:
        return float(s)
    except ValueError:
        return float("nan")


def _ffill(row: tuple) -> list:
    out, last = [], None
    for v in row:
        text = _clean_text(v)
        if text:
            last = text
        out.append(last)
    return out


@dataclass
class ColumnSpec:
    index: int
    metric: str
    instrument: str | None
    unit: str


def _classify_columns(row2, row3, row4, row5, row6) -> list[ColumnSpec]:
    specs = []
    for i, n in enumerate(row6):
        if not isinstance(n, (int, float)):
            continue  # only numbered metric columns (D onward) are data columns
        g2 = (row2[i] or "").lower()
        g5 = row5[i] or ""
        g5l = g5.lower()

        if "outstanding" in g2:
            metric_base, unit = "outstanding_count", "count"
        elif "active" in g2:
            metric_base, unit = "active_count", "count"
        elif "cash withdrawal" in g2:
            metric_base, unit = None, None  # resolved below via row4
        elif "payment transactions" in g2:
            metric_base, unit = None, None
        else:
            continue

        if metric_base in ("outstanding_count", "active_count"):
            instrument = "card" if "cards" in g5l else "wallet" if "wallets" in g5l else None
            specs.append(ColumnSpec(i, metric_base, instrument, unit))
            continue

        g3 = (row3[i] or "").lower()
        g4 = (row4[i] or "").lower()
        instrument = "card" if "cards" in g3 else "wallet" if "wallets" in g3 else None
        if "purchase" in g4:
            subtype = "purchase"
        elif "fund transfer" in g4:
            subtype = "fund_transfer"
        elif "atm" in g4:
            subtype = "cash_withdrawal_atm"
        elif "pos" in g4:
            subtype = "cash_withdrawal_pos"
        else:
            continue
        value_type = "volume" if "volume" in g5l else "value" if "value" in g5l else None
        if value_type is None or instrument is None:
            continue
        unit = "count" if value_type == "volume" else "inr_thousand"
        specs.append(ColumnSpec(i, f"{subtype}_txn_{value_type}", instrument, unit))
    return specs


def load_aliases() -> dict[str, str]:
    if not ALIASES_PATH.exists():
        ALIASES_PATH.parent.mkdir(parents=True, exist_ok=True)
        ALIASES_PATH.write_text("raw_name,canonical_name\n", encoding="utf-8")
        return {}
    with open(ALIASES_PATH, encoding="utf-8") as f:
        return {row["raw_name"]: row["canonical_name"] for row in csv.DictReader(f)}


def _append_identity_aliases(new_names: list[str]) -> None:
    """Register newly-seen raw names as identity rows (raw==canonical) so the alias
    file becomes an editable registry; never merges names on its own."""
    if not new_names:
        return
    with open(ALIASES_PATH, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for name in new_names:
            writer.writerow([name, name])


def record_renames(renames: dict[str, str]) -> None:
    """Point each renamed issuer's alias row at its current name, so the merge applied
    during refresh is visible and editable here rather than being invisible magic."""
    if not renames:
        return
    aliases = load_aliases()
    aliases.update(renames)
    with open(ALIASES_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["raw_name", "canonical_name"])
        for raw, canonical in sorted(aliases.items()):
            writer.writerow([raw, canonical])


def parse_file(path: Path, month: date, is_revised: bool) -> tuple[pd.DataFrame, list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, max_row=6, values_only=True))
    row2, row3, row4, row5, row6 = (_ffill(r) if i != 4 else list(r) for i, r in enumerate(rows))
    specs = _classify_columns(row2, row3, row4, row5, row6)

    aliases = load_aliases()
    unmatched: list[str] = []
    records = []
    section = None
    for row in ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=True):
        b, c = row[1], row[2]
        if isinstance(b, str):
            bl = b.strip().lower()
            if bl in ("banks",):
                section = "bank"
            elif bl in ("non banks", "non-banks", "non bank", "non-bank"):
                section = "non_bank"
            elif bl.startswith("total"):
                break
            continue
        if not isinstance(b, (int, float)) or not c:
            continue
        entity_raw = _clean_text(c)
        if not entity_raw or section is None:
            continue
        canonical = aliases.get(entity_raw)
        if canonical is None:
            canonical = entity_raw
            unmatched.append(entity_raw)

        for spec in specs:
            records.append({
                "month": month,
                "entity_raw": entity_raw,
                "entity": canonical,
                "entity_type": section,
                "instrument": spec.instrument,
                "metric": spec.metric,
                "unit": spec.unit,
                "value": coerce_number(row[spec.index]),
                "source_file": path.name,
                "is_revised": is_revised,
            })

    new_names = sorted(set(unmatched))
    _append_identity_aliases(new_names)
    return pd.DataFrame.from_records(records), new_names
