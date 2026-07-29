"""Phase 0: dump the real structure of sample XLSX files. Read-only, no parsing logic."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

RAW_DIR = Path(__file__).resolve().parent.parent / "data" / "raw"


def inspect(path: Path) -> None:
    print(f"\n{'=' * 100}\n{path.name}\n{'=' * 100}")
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- sheet: {sheet_name!r}  dims={ws.dimensions}  ({ws.max_row} rows x {ws.max_column} cols) ---")
        print("merged ranges:", [str(r) for r in ws.merged_cells.ranges])
        print("\nfirst 15 rows:")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            print(f"  row {i}: {row}")


if __name__ == "__main__":
    files = sys.argv[1:] or sorted(RAW_DIR.glob("*.xlsx"))
    for f in files:
        inspect(Path(f))
